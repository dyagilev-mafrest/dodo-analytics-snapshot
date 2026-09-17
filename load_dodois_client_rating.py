# Loads the manually-exported Dodo IS «Рейтинг клиентов» (виджет `Client Rating
# Overall Table`) into dodois_client_rating — эталон рядом с нашим расчётом из
# customer-feedback/customer-ratings. See migrations/056_dodois_client_rating.sql
# и docs/customer-rating-methodology.md.
#
# Not an API pull — no scheduled job; re-run manually на свежей выгрузке.
# ON CONFLICT перезаписывает уже загруженный период.
#
# Usage:
#   python load_dodois_client_rating.py export.xlsx --exported-at 2026-09-02
#   python load_dodois_client_rating.py export.xlsx --exported-at 2026-09-02 --period-days 7
#
# Колонки ищутся по заголовку (регистр и пробелы не важны, достаточно вхождения
# подстроки, выигрывает самый длинный алиас). Доли из выгрузки не храним — они
# ровно равны отношениям сохранённых счётчиков (проверено на выгрузке 02.09.2026).
import csv
from datetime import date, datetime, timedelta

import psycopg2.extras

from db import get_connection, get_cursor

COLUMN_ALIASES = {
    "period":                   ["период", "неделя", "week", "месяц", "month", "дата"],
    "unit_name":                ["пиццерия", "unit", "подразделение"],
    "rating_total":             ["рк общий"],
    "rating_delivery":          ["рк доставка"],
    "rating_restaurant":        ["рк ресторан"],
    "ratings_count":            ["кол-во оценок"],
    "ratings_count_delivery":   ["кол-во оценок доставка"],
    "ratings_count_restaurant": ["кол-во оценок ресторан"],
    "delivery_orders_app":      ["кол-во заказов на доставку через приложение"],
    "delivery_orders_all":      ["кол-во всех заказов на доставку"],
    "restaurant_orders_all":    ["кол-во всех заказов в ресторан"],
    "restaurant_orders_app":    ["кол-во заказов в ресторан через приложение"],
}

RATING_FIELDS = ("rating_total", "rating_delivery", "rating_restaurant")
COUNT_FIELDS = ("ratings_count", "ratings_count_delivery", "ratings_count_restaurant",
                "delivery_orders_app", "delivery_orders_all",
                "restaurant_orders_all", "restaurant_orders_app")


def norm(s) -> str:
    return " ".join(str(s or "").lower().replace("\xa0", " ").split())


def map_headers(headers: list) -> tuple[dict, list]:
    """(индекс колонки -> поле, нераспознанные заголовки). Самый длинный алиас
    выигрывает: «кол-во оценок доставка» не должна достаться «кол-во оценок»."""
    mapping, unknown = {}, []
    for i, h in enumerate(headers):
        h_norm = norm(h)
        if not h_norm:
            continue
        best_field, best_len = None, 0
        for field, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias in h_norm and len(alias) > best_len:
                    best_field, best_len = field, len(alias)
        if best_field:
            mapping[i] = best_field
        else:
            unknown.append(str(h))
    return mapping, unknown


def to_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace("%", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def to_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def read_table(path: str, sheet: str | None) -> tuple[list, list[list]]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            rows = [r for r in csv.reader(f, csv.Sniffer().sniff(sample, delimiters=";,\t"))]
    for i, row in enumerate(rows):
        mapping, _ = map_headers(row)
        if "unit_name" in mapping.values():
            return row, rows[i + 1:]
    raise SystemExit("Не нашёл строку заголовка с колонкой «Пиццерия» — проверьте файл/лист.")


def period_ends(starts: set, period_days: int | None) -> dict:
    """Выгрузка помечает строку датой НАЧАЛА периода. Если периодов несколько —
    длину берём из шага между ними, если один — из --period-days."""
    starts = sorted(starts)
    if len(starts) >= 2:
        step = min((b - a).days for a, b in zip(starts, starts[1:]))
        ends = {s: starts[i + 1] - timedelta(days=1) for i, s in enumerate(starts[:-1])}
        ends[starts[-1]] = starts[-1] + timedelta(days=step - 1)
        return ends
    days = period_days or 7
    print(f"В выгрузке один период — считаю его длиной {days} дн. "
          f"(меняется через --period-days).")
    return {starts[0]: starts[0] + timedelta(days=days - 1)}


def parse_rows(path: str, sheet: str | None, exported_at: date, period_days: int | None) -> list[dict]:
    headers, body = read_table(path, sheet)
    mapping, unknown = map_headers(headers)
    if unknown:
        print(f"Нераспознанные колонки (игнорирую): {', '.join(unknown)}")

    raw = []
    for row in body:
        rec = {field: (row[i] if i < len(row) else None) for i, field in mapping.items()}
        if not norm(rec.get("unit_name")):
            continue
        start = to_date(rec.get("period"))
        if start is None:
            raise SystemExit("Не разобрал колонку периода.")
        rec["_start"] = start
        raw.append(rec)
    if not raw:
        raise SystemExit("В файле нет строк с пиццериями.")

    ends = period_ends({r["_start"] for r in raw}, period_days)
    out = []
    for rec in raw:
        item = {
            "period_start": rec["_start"],
            "period_end": ends[rec["_start"]],
            "unit_name": str(rec["unit_name"]).strip(),
            "exported_at": exported_at,
        }
        for f in RATING_FIELDS:
            item[f] = to_number(rec.get(f))
        for f in COUNT_FIELDS:
            v = to_number(rec.get(f))
            item[f] = int(v) if v is not None else None
        # Выгрузка ставит 0 там, где канала нет вовсе (у Якутск-5 нет зала, значит
        # «РК ресторан» = 0 при нуле оценок). На шкале 1–5 ноль не оценка, а «нет
        # данных» — иначе он поедет в средние и в таблицы по пиццериям как реальные
        # ноль баллов.
        if not item["ratings_count_delivery"]:
            item["rating_delivery"] = None
        if not item["ratings_count_restaurant"]:
            item["rating_restaurant"] = None
        if not item["ratings_count"]:
            item["rating_total"] = None
        out.append(item)
    return out


COLUMNS = ("period_start", "period_end", "unit_name", *RATING_FIELDS, *COUNT_FIELDS, "exported_at")


def upsert_rows(rows: list[dict]):
    conn = get_connection()
    cur = get_cursor(conn)
    updates = ", ".join(f"{c} = EXCLUDED.{c}" for c in COLUMNS[3:])
    psycopg2.extras.execute_values(
        cur,
        f"""
        INSERT INTO dodois_client_rating ({", ".join(COLUMNS)})
        VALUES %s
        ON CONFLICT (period_start, period_end, unit_name) DO UPDATE SET
            {updates}, loaded_at = now()
        """,
        [tuple(r[c] for c in COLUMNS) for r in rows],
        page_size=500,
    )
    conn.commit()
    conn.close()


def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("path", help="xlsx или csv выгрузка виджета «Client Rating Overall Table»")
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--exported-at", required=True, help="дата выгрузки из Dodo IS (YYYY-MM-DD)")
    parser.add_argument("--period-days", type=int, default=None,
                        help="длина периода, если в выгрузке он один (по умолчанию 7)")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    rows = parse_rows(args.path, args.sheet, date.fromisoformat(args.exported_at), args.period_days)
    for r in sorted(rows, key=lambda r: (r["period_start"], r["unit_name"])):
        print(f"  {r['period_start']}..{r['period_end']} | {r['unit_name']:<10} | "
              f"доставка {r['rating_delivery'] or 0:.4f} ({r['ratings_count_delivery']}) | "
              f"ресторан {r['rating_restaurant'] or 0:.4f} ({r['ratings_count_restaurant']})")
    if args.dry_run:
        print(f"\n--dry-run: {len(rows)} строк не записано.")
        return
    upsert_rows(rows)
    print(f"\n💾 Загружено строк: {len(rows)}")


if __name__ == "__main__":
    main()
