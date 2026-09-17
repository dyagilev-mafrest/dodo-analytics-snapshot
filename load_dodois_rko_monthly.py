# ⚠️ ПУТЬ ЗАКРЫТ 11.09.2026 — не запускать. Заменён на get_ratings.py: живой
# pull из controlling-api, по ПИЦЦЕРИЯМ, в ratings_unit_period (см.
# migrations/060). Причина, записанная ниже — «нет нужного scope» — оказалась
# неверной: 401 был протухшим токеном, а не отказом в правах. С живым токеном
# оба эндпоинта отдают 200, а цифры совпадают с дашбордом Dodo IS до сотых.
#
# Таблица dodois_rko_monthly НЕ удалена намеренно: это единственная история РКО
# до 11.09.2026. По API прошлые периоды достаются только через
# /ratings/customer-experience/history, а по пиццериям их нет вообще — история
# копится с первых прогонов нового загрузчика. Читать таблицу можно, дописывать
# в неё — нет.
#
# Loads the manually-exported Dodo IS Superset xlsx pair (РКО — рейтинг
# клиентского опыта, endpoint controlling/ratings/customer-experience) into
# dodois_rko_monthly. See migrations/043_dodois_rko_monthly.sql for why this
# is a manual export rather than a live API pull.
#
# Not an API pull — no scheduled job; re-run manually whenever a fresh pair
# of exports is provided. Only current-year data is kept (see caller/loader
# convention elsewhere in this repo) — table stays tiny by design.
#
# Usage:
#   python load_dodois_rko_monthly.py <rko_and_store_number.xlsx> <rko_last12.xlsx>
#
# File 1 columns:  Начало проверки | Кол-во пиццерий | Средний РКО
# File 2 columns:  Начало проверки | Кол-во пиццерий 12 последних проверок | Средний РКО (12 проверок)
import sys

import openpyxl

from db import get_connection, get_cursor


def load_main(path: str, sheet: str = "Sheet1") -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    out = {}
    for period, count, avg_rko in ws.iter_rows(min_row=2, values_only=True):
        if period is None:
            continue
        month = period.date() if hasattr(period, "date") else period
        out[month] = {"pizzeria_count": count, "avg_rko": avg_rko}
    return out


def load_last12(path: str, sheet: str = "Sheet1") -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    out = {}
    for period, count, avg_rko in ws.iter_rows(min_row=2, values_only=True):
        if period is None:
            continue
        month = period.date() if hasattr(period, "date") else period
        out[month] = {"last12_pizzeria_count": count, "avg_rko_last12": avg_rko}
    return out


def upsert_rows(main: dict, last12: dict):
    months = sorted(set(main) | set(last12))
    rows = []
    for month in months:
        m = main.get(month, {})
        l12 = last12.get(month, {})
        rows.append((
            month,
            m.get("pizzeria_count"), m.get("avg_rko"),
            l12.get("last12_pizzeria_count"), l12.get("avg_rko_last12"),
        ))
    if not rows:
        return
    import psycopg2.extras
    conn = get_connection()
    cur = get_cursor(conn)
    psycopg2.extras.execute_values(cur, """
        INSERT INTO dodois_rko_monthly
            (month, pizzeria_count, avg_rko, last12_pizzeria_count, avg_rko_last12, loaded_at)
        VALUES %s
        ON CONFLICT (month) DO UPDATE SET
            pizzeria_count        = EXCLUDED.pizzeria_count,
            avg_rko               = EXCLUDED.avg_rko,
            last12_pizzeria_count = EXCLUDED.last12_pizzeria_count,
            avg_rko_last12        = EXCLUDED.avg_rko_last12,
            loaded_at             = now()
    """, rows, template="(%s,%s,%s,%s,%s,now())", page_size=200)
    conn.commit()
    conn.close()
    print(f"Loaded {len(rows)} months: {months[0]} .. {months[-1]}")


if __name__ == "__main__":
    main_path, last12_path = sys.argv[1], sys.argv[2]
    main = load_main(main_path)
    last12 = load_last12(last12_path)
    upsert_rows(main, last12)
