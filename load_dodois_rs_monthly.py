# ⚠️ ПУТЬ ЗАКРЫТ 11.09.2026 — не запускать. Заменён на get_ratings.py: живой
# pull из controlling-api (/ratings/standards — не customer-experience, как
# ошибочно сказано ниже), по ПИЦЦЕРИЯМ, в ratings_unit_period (см.
# migrations/060).
#
# Таблица dodois_rs_monthly НЕ удалена намеренно: это единственная история РС
# до 11.09.2026. По API прошлые периоды достаются только через
# /ratings/standards/history, а по пиццериям их нет вообще — история копится с
# первых прогонов нового загрузчика. Читать таблицу можно, дописывать — нет.
#
# Loads the manually-exported Dodo IS Superset xlsx pair (РС — рейтинг
# стандартов, endpoint controlling/ratings/customer-experience) into
# dodois_rs_monthly. See migrations/044_dodois_rs_monthly.sql for why this
# is a manual export rather than a live API pull. Same shape/pattern as
# load_dodois_rko_monthly.py, just a different metric (РС, not РКО) and a
# 6-check rolling window instead of 12.
#
# Not an API pull — no scheduled job; re-run manually whenever a fresh pair
# of exports is provided. Only current-year data is kept — table stays tiny.
#
# Usage:
#   python load_dodois_rs_monthly.py <rs_and_store_number.xlsx> <rs_last6.xlsx>
#
# File 1 columns:  Начало проверки | Кол-во пиццерий | Средний РC
# File 2 columns:  Начало проверки | Кол-во пиццерий 6 последних проверок | Средний РC (6 проверок)
import sys

import openpyxl

from db import get_connection, get_cursor


def load_main(path: str, sheet: str = "Sheet1") -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    out = {}
    for period, count, avg_rs in ws.iter_rows(min_row=2, values_only=True):
        if period is None:
            continue
        month = period.date() if hasattr(period, "date") else period
        out[month] = {"pizzeria_count": count, "avg_rs": avg_rs}
    return out


def load_last6(path: str, sheet: str = "Sheet1") -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    out = {}
    for period, count, avg_rs in ws.iter_rows(min_row=2, values_only=True):
        if period is None:
            continue
        month = period.date() if hasattr(period, "date") else period
        out[month] = {"last6_pizzeria_count": count, "avg_rs_last6": avg_rs}
    return out


def upsert_rows(main: dict, last6: dict):
    months = sorted(set(main) | set(last6))
    rows = []
    for month in months:
        m = main.get(month, {})
        l6 = last6.get(month, {})
        rows.append((
            month,
            m.get("pizzeria_count"), m.get("avg_rs"),
            l6.get("last6_pizzeria_count"), l6.get("avg_rs_last6"),
        ))
    if not rows:
        return
    import psycopg2.extras
    conn = get_connection()
    cur = get_cursor(conn)
    psycopg2.extras.execute_values(cur, """
        INSERT INTO dodois_rs_monthly
            (month, pizzeria_count, avg_rs, last6_pizzeria_count, avg_rs_last6, loaded_at)
        VALUES %s
        ON CONFLICT (month) DO UPDATE SET
            pizzeria_count       = EXCLUDED.pizzeria_count,
            avg_rs               = EXCLUDED.avg_rs,
            last6_pizzeria_count = EXCLUDED.last6_pizzeria_count,
            avg_rs_last6         = EXCLUDED.avg_rs_last6,
            loaded_at            = now()
    """, rows, template="(%s,%s,%s,%s,%s,now())", page_size=200)
    conn.commit()
    conn.close()
    print(f"Loaded {len(rows)} months: {months[0]} .. {months[-1]}")


if __name__ == "__main__":
    main_path, last6_path = sys.argv[1], sys.argv[2]
    main = load_main(main_path)
    last6 = load_last6(last6_path)
    upsert_rows(main, last6)
